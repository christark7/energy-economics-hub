import runpy
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from export import export_energy_hub_data
from market_intelligence import build_market_rankings, generate_company_market_data
from scripts import ingest
from scripts.ingest import generate_energy_data
from scripts.model_engine import run_model
from scripts.vif_check import calculate_vif
from scripts.web_data_ingestion import (
    build_cross_source_spreads,
    fetch_configured_sources,
    fetch_commodities_api_latest,
    fetch_eia_petroleum_prices,
    fetch_http_market_data,
    normalize_market_records,
)
from stakeholder_packager import generate_stakeholder_package

import pytest

@pytest.fixture(autouse=True)
def setup_model_outputs():
    """Ensure output CSVs exist before any test runs."""
    Path("outputs").mkdir(exist_ok=True)
    run_model()


def test_generate_energy_data_has_expected_columns():
    df = generate_energy_data()

    expected_columns = {
        "timestamp",
        "region",
        "load_mw",
        "spot_price_mwh",
        "margin_estimate_pct",
    }

    assert expected_columns.issubset(df.columns)
    assert len(df) == 24

def test_run_model_returns_valid_metrics():
    result = run_model()

    assert "model" in result
    assert "coefficients" in result
    assert "metrics" in result
    assert "R-Squared" in result["metrics"]["Metric"].values


def test_calculate_vif_returns_positive_values():
    df = generate_energy_data()
    vif_df = calculate_vif(df)

    assert {"Variable", "VIF"}.issubset(vif_df.columns)
    assert (vif_df["VIF"] >= 1).all()


def test_export_energy_hub_data_creates_workbook():
    output_path = export_energy_hub_data()
    workbook = load_workbook(output_path)

    assert set(workbook.sheetnames) == {"Coefficients", "Model_Metrics", "VIF_Diagnostics"}
    assert Path(output_path).exists()


def test_generate_stakeholder_package_creates_business_files():
    package_paths = generate_stakeholder_package()

    assert all(Path(path).exists() for path in package_paths.values())
    assert {"xlsx", "docx"}.issubset({path.suffix.lstrip(".") for path in package_paths.values()})


def test_market_intelligence_returns_ranked_companies():
    companies = generate_company_market_data()
    rankings = build_market_rankings(companies)

    assert len(companies) == 10
    assert len(rankings) == 10
    assert rankings["profitability_score"].is_monotonic_decreasing
    assert {"company_name", "sector", "service", "profitability_score", "arbitrage_score"}.issubset(rankings.columns)


def test_blob_upload_uses_default_azure_credential(monkeypatch, tmp_path):
    uploaded = {}

    class FakeCredential:
        def close(self):
            uploaded["credential_closed"] = True

    class FakeBlobClient:
        def upload_blob(self, data, overwrite):
            uploaded["content"] = data.read()
            uploaded["overwrite"] = overwrite

    class FakeBlobServiceClient:
        def __init__(self, account_url, credential):
            uploaded["account_url"] = account_url
            uploaded["credential"] = credential

        def get_blob_client(self, container, blob):
            uploaded["container"] = container
            uploaded["blob"] = blob
            return FakeBlobClient()

        def close(self):
            uploaded["client_closed"] = True

    source = tmp_path / "sample.csv"
    source.write_bytes(b"timestamp,value\n2026-10-10,1\n")
    monkeypatch.delenv("AZURE_STORAGE_CONNECTION_STRING", raising=False)
    monkeypatch.setenv("AZURE_STORAGE_CONTAINER", "energyblob")
    monkeypatch.setattr(ingest, "DefaultAzureCredential", FakeCredential)
    monkeypatch.setattr(ingest, "BlobServiceClient", FakeBlobServiceClient)

    blob_name = ingest.upload_to_blob_storage(str(source))

    assert blob_name is not None
    assert uploaded["account_url"] == "https://energystoragechris.blob.core.windows.net"
    assert uploaded["container"] == "energyblob"
    assert uploaded["content"] == source.read_bytes()
    assert uploaded["overwrite"] is True
    assert uploaded["credential_closed"] is True
    assert uploaded["client_closed"] is True


def test_blob_upload_uses_requested_blob_name(monkeypatch, tmp_path):
    uploaded = {}

    class FakeBlobClient:
        def upload_blob(self, data, overwrite):
            uploaded["content"] = data.read()
            uploaded["overwrite"] = overwrite

    class FakeBlobServiceClient:
        def __init__(self, account_url, credential):
            pass

        def get_blob_client(self, container, blob):
            uploaded["container"] = container
            uploaded["blob"] = blob
            return FakeBlobClient()

        def close(self):
            pass

    class FakeCredential:
        def close(self):
            pass

    source = tmp_path / "market_observations.csv"
    source.write_text("observed_at,price\n", encoding="utf-8")
    monkeypatch.delenv("AZURE_STORAGE_CONNECTION_STRING", raising=False)
    monkeypatch.setattr(ingest, "DefaultAzureCredential", FakeCredential)
    monkeypatch.setattr(ingest, "BlobServiceClient", FakeBlobServiceClient)

    uploaded_blob = ingest.upload_to_blob_storage(
        str(source),
        container_name="energyblob",
        blob_name="market-data/market_observations.csv",
    )

    assert uploaded_blob == "market-data/market_observations.csv"
    assert uploaded["container"] == "energyblob"
    assert uploaded["blob"] == "market-data/market_observations.csv"


def test_normalize_market_records_preserves_source_provenance():
    records = [
        {"date": "2026-10-10T12:00:00-04:00", "value": "42.5", "zone": "ZONE-A"},
    ]

    normalized = normalize_market_records(
        records,
        "Feed A",
        {"observed_at": "date", "price": "value", "region": "zone"},
        {"commodity": "electricity", "currency": "USD", "unit": "MWh"},
    )

    assert normalized.iloc[0]["observed_at"] == pd.Timestamp("2026-10-10T16:00:00Z")
    assert normalized.iloc[0]["source"] == "Feed A"
    assert normalized.iloc[0]["price"] == 42.5
    assert normalized.iloc[0]["region"] == "ZONE-A"


def test_cross_source_spreads_only_compare_matching_market_dimensions():
    records = [
        {"observed_at": "2026-10-10T12:00:00Z", "source": "Feed A", "market": "spot", "region": "ZONE-A", "commodity": "electricity", "price": 42.0, "currency": "USD", "unit": "MWh", "company": None},
        {"observed_at": "2026-10-10T12:00:00Z", "source": "Feed B", "market": "spot", "region": "ZONE-A", "commodity": "electricity", "price": 49.0, "currency": "USD", "unit": "MWh", "company": None},
        {"observed_at": "2026-10-10T12:00:00Z", "source": "Feed C", "market": "spot", "region": "ZONE-B", "commodity": "electricity", "price": 90.0, "currency": "USD", "unit": "MWh", "company": None},
        {"observed_at": "2026-10-10T12:00:00Z", "source": "Feed D", "market": "retail", "region": "ZONE-A", "commodity": "electricity", "price": 0.08, "currency": "USD", "unit": "kWh", "company": None},
        {"observed_at": "2026-10-10T12:00:00Z", "source": "Feed E", "market": "day-ahead", "region": "ZONE-A", "commodity": "electricity", "price": 90.0, "currency": "USD", "unit": "MWh", "company": None},
    ]
    observations = pd.DataFrame(records)

    spreads = build_cross_source_spreads(observations)

    assert len(spreads) == 1
    assert spreads.iloc[0]["buy_from_source"] == "Feed A"
    assert spreads.iloc[0]["sell_to_source"] == "Feed B"
    assert spreads.iloc[0]["gross_spread"] == 7.0
    assert spreads.iloc[0]["gross_spread_pct"] == (7.0 / 42.0) * 100


def test_cross_source_spreads_can_compare_companies_from_one_feed():
    observations = pd.DataFrame([
        {"observed_at": "2026-10-10T12:00:00Z", "source": "Supplier API", "market": "retail", "region": "ZONE-A", "commodity": "petrol", "price": 1.50, "currency": "USD", "unit": "litre", "company": "Supplier One"},
        {"observed_at": "2026-10-10T12:00:00Z", "source": "Supplier API", "market": "retail", "region": "ZONE-A", "commodity": "petrol", "price": 1.65, "currency": "USD", "unit": "litre", "company": "Supplier Two"},
    ])

    spreads = build_cross_source_spreads(observations)

    assert len(spreads) == 1
    assert spreads.iloc[0]["buy_company"] == "Supplier One"
    assert spreads.iloc[0]["sell_company"] == "Supplier Two"
    assert spreads.iloc[0]["source_count"] == 1
    assert spreads.iloc[0]["quote_count"] == 2
    assert round(spreads.iloc[0]["gross_spread"], 2) == 0.15


def test_cross_source_spreads_exclude_stale_and_synthetic_prices():
    observations = pd.DataFrame([
        {"observed_at": "2026-10-10T12:00:00Z", "source": "Current A", "market": "spot", "region": "global", "commodity": "brent", "price": 70.0, "currency": "USD", "unit": "barrel", "company": None, "is_stale": False, "is_synthetic": False},
        {"observed_at": "2026-10-10T12:00:00Z", "source": "Stale B", "market": "spot", "region": "global", "commodity": "brent", "price": 80.0, "currency": "USD", "unit": "barrel", "company": None, "is_stale": True, "is_synthetic": False},
        {"observed_at": "2026-10-10T12:00:00Z", "source": "Synthetic C", "market": "spot", "region": "global", "commodity": "brent", "price": 90.0, "currency": "USD", "unit": "barrel", "company": None, "is_stale": False, "is_synthetic": True},
        {"observed_at": "2026-10-10T12:00:00Z", "source": "Current D", "market": "spot", "region": "global", "commodity": "brent", "price": 72.0, "currency": "USD", "unit": "barrel", "company": None, "is_stale": False, "is_synthetic": False},
    ])

    spreads = build_cross_source_spreads(observations)

    assert len(spreads) == 1
    assert spreads.iloc[0]["buy_from_source"] == "Current A"
    assert spreads.iloc[0]["sell_to_source"] == "Current D"
    assert spreads.iloc[0]["gross_spread"] == 2.0


def test_eia_petroleum_fetch_uses_weekly_retail_facets(monkeypatch):
    request_options = {}

    class FakeResponse:
        def raise_for_status(self):
            pass

        def json(self):
            return {"response": {"data": [{
                "period": "2026-10-05",
                "duoarea": "NUS",
                "area-name": "U.S.",
                "product": "EPMR",
                "product-name": "Regular Gasoline",
                "process": "PTE",
                "value": "3.25",
                "units": "$/GAL",
                "series-description": "U.S. regular gasoline retail prices",
            }]}}

    def fake_get(url, **kwargs):
        request_options.update(kwargs)
        return FakeResponse()

    monkeypatch.setattr("scripts.web_data_ingestion.requests.get", fake_get)

    prices = fetch_eia_petroleum_prices("test-key", "EPMR", "NUS", "2026-01-01", "2026-10-10")

    assert len(prices) == 1
    assert prices.iloc[0]["commodity"] == "gasoline_regular"
    assert prices.iloc[0]["unit"] == "$/GAL"
    assert request_options["params"]["frequency"] == "weekly"
    assert request_options["params"]["facets[product][]"] == "EPMR"
    assert request_options["params"]["facets[duoarea][]"] == "NUS"
    assert request_options["params"]["facets[process][]"] == "PTE"


def test_commodities_api_inverts_inverse_rates(monkeypatch):
    request_options = {}

    class FakeResponse:
        def raise_for_status(self):
            pass

        def json(self):
            return {
                "success": True,
                "timestamp": 1791633600,
                "base": "USD",
                "rates": {"WTIOIL": 0.05, "BRENTOIL": 0.04},
                "unit": {"WTIOIL": "per barrel", "BRENTOIL": "per barrel"},
            }

    def fake_get(url, **kwargs):
        request_options.update(kwargs)
        return FakeResponse()

    monkeypatch.setattr("scripts.web_data_ingestion.requests.get", fake_get)

    prices = fetch_commodities_api_latest("test-key", ["WTIOIL", "BRENTOIL"])

    assert prices.set_index("commodity").loc["WTIOIL", "price"] == 20.0
    assert prices.set_index("commodity").loc["BRENTOIL", "price"] == 25.0
    assert request_options["params"]["access_key"] == "test-key"
    assert request_options["params"]["symbols"] == "WTIOIL,BRENTOIL"


def test_http_feed_maps_json_records(monkeypatch):
    class FakeResponse:
        text = ""

        def raise_for_status(self):
            pass

        def json(self):
            return {"data": [{"time": "2026-10-10T12:00:00Z", "usd_mwh": "55.2", "zone": "ZONE-A"}]}

    monkeypatch.setattr("scripts.web_data_ingestion.requests.get", lambda *args, **kwargs: FakeResponse())
    source = {
        "name": "Supplier feed",
        "url": "https://provider.example/prices",
        "records_path": "data",
        "field_map": {"observed_at": "time", "price": "usd_mwh", "region": "zone"},
        "defaults": {"market": "day-ahead", "commodity": "electricity", "currency": "USD", "unit": "MWh"},
    }

    observations = fetch_http_market_data(source)

    assert len(observations) == 1
    assert observations.iloc[0]["source"] == "Supplier feed"
    assert observations.iloc[0]["price"] == 55.2
    assert observations.iloc[0]["unit"] == "MWh"


def test_oilpriceapi_feed_uses_token_auth_and_preserves_freshness(monkeypatch):
    request_options = {}

    class FakeResponse:
        text = ""

        def raise_for_status(self):
            pass

        def json(self):
            return {"data": {"prices": [{
                "as_of": "2026-10-10T12:00:00Z",
                "created_at": "2026-10-10T12:03:00Z",
                "price": 75.0,
                "code": "BRENT_CRUDE_USD",
                "currency": "USD",
                "unit": "barrel",
                "type": "spot_price",
                "stale": False,
                "synthetic": False,
                "source": "market_reporting",
            }]}}

    def fake_get(url, **kwargs):
        request_options.update(kwargs)
        return FakeResponse()

    monkeypatch.setenv("OILPRICEAPI_TEST_KEY", "test-token")
    monkeypatch.setattr("scripts.web_data_ingestion.requests.get", fake_get)
    source = {
        "name": "OilPriceAPI",
        "url": "https://api.oilpriceapi.com/v1/prices/latest",
        "format": "json",
        "records_path": "data.prices",
        "token_env": "OILPRICEAPI_TEST_KEY",
        "auth_prefix": "Token ",
        "field_map": {
            "observed_at": ["as_of", "created_at"],
            "price": "price",
            "commodity": "code",
            "currency": "currency",
            "unit": "unit",
            "market": "type",
            "source_detail": "source",
            "is_stale": "stale",
            "is_synthetic": "synthetic",
        },
        "defaults": {"region": "global benchmark"},
    }

    observations = fetch_http_market_data(source)

    assert observations.iloc[0]["observed_at"] == pd.Timestamp("2026-10-10T12:00:00Z")
    assert observations.iloc[0]["source_detail"] == "market_reporting"
    assert observations.iloc[0]["is_stale"] == False
    assert observations.iloc[0]["is_synthetic"] == False
    assert request_options["headers"]["Authorization"] == "Token test-token"


def test_http_xml_feed_maps_records_with_basic_auth(monkeypatch):
    request_options = {}

    class FakeResponse:
        content = b"""<feed><item><date>2026-10-10T12:00:00Z</date><price>1.72</price><country>US</country><fuel>gasoline</fuel><currency>USD</currency><unit>litre</unit></item></feed>"""

        def raise_for_status(self):
            pass

    def fake_get(url, **kwargs):
        request_options.update(kwargs)
        return FakeResponse()

    monkeypatch.setenv("FUEL_FEED_USERNAME", "test-user")
    monkeypatch.setenv("FUEL_FEED_PASSWORD", "test-password")
    monkeypatch.setattr("scripts.web_data_ingestion.requests.get", fake_get)
    source = {
        "name": "Licensed fuel feed",
        "url": "https://provider.example/feed.xml",
        "format": "xml",
        "record_tag": "item",
        "username_env": "FUEL_FEED_USERNAME",
        "password_env": "FUEL_FEED_PASSWORD",
        "field_map": {
            "observed_at": "date",
            "price": "price",
            "region": "country",
            "commodity": "fuel",
            "currency": "currency",
            "unit": "unit",
        },
        "defaults": {"market": "retail"},
    }

    observations = fetch_http_market_data(source)

    assert len(observations) == 1
    assert observations.iloc[0]["commodity"] == "gasoline"
    assert observations.iloc[0]["region"] == "US"
    assert request_options["auth"] == ("test-user", "test-password")


def test_unconfigured_feed_credentials_are_reported_without_crashing(monkeypatch):
    monkeypatch.delenv("TEST_FEED_TOKEN", raising=False)

    observations, errors = fetch_configured_sources([
        {
            "type": "http",
            "name": "Supplier feed",
            "url": "https://provider.example/prices",
            "token_env": "TEST_FEED_TOKEN",
        }
    ])

    assert observations.empty
    assert len(errors) == 1
    assert "TEST_FEED_TOKEN" in errors[0]


def test_disabled_example_feed_is_skipped():
    observations, errors = fetch_configured_sources([
        {"type": "http", "name": "Disabled placeholder", "enabled": False}
    ])

    assert observations.empty
    assert errors == []


def test_configured_market_outputs_upload_to_selected_container(monkeypatch, tmp_path):
    from scripts import web_data_ingestion

    config_path = tmp_path / "feeds.json"
    config_path.write_text('{"sources": []}', encoding="utf-8")
    output_dir = tmp_path / "outputs"
    uploaded = []
    observations = pd.DataFrame([{
        "observed_at": pd.Timestamp("2026-10-10T12:00:00Z"),
        "source": "Test feed",
        "market": "retail",
        "region": "US",
        "commodity": "gasoline",
        "price": 3.25,
        "currency": "USD",
        "unit": "gallon",
        "company": None,
        "source_detail": None,
        "is_stale": pd.NA,
        "is_synthetic": pd.NA,
    }])

    monkeypatch.setenv("AZURE_STORAGE_UPLOADS", "true")
    monkeypatch.setenv("AZURE_STORAGE_CONTAINER", "energyblob")
    monkeypatch.setattr(web_data_ingestion, "load_feed_config", lambda path: [])
    monkeypatch.setattr(web_data_ingestion, "fetch_configured_sources", lambda sources: (observations, []))
    monkeypatch.setattr(
        ingest,
        "upload_to_blob_storage",
        lambda file_path, container_name, blob_name: uploaded.append((Path(file_path).name, container_name, blob_name)) or blob_name,
    )

    observation_path, spread_path = web_data_ingestion.run_configured_ingestion(config_path, output_dir)

    assert observation_path.exists()
    assert spread_path.exists()
    assert {(item[0], item[1], item[2]) for item in uploaded} == {
        ("market_observations.csv", "energyblob", "market-data/market_observations.csv"),
        ("market_spread_opportunities.csv", "energyblob", "market-data/market_spread_opportunities.csv"),
    }


def test_market_ingestion_direct_script_adds_repository_import_path(monkeypatch, tmp_path):
    script_path = Path(__file__).resolve().parents[1] / "scripts" / "web_data_ingestion.py"
    project_root = script_path.parents[1]
    config_path = tmp_path / "empty_feeds.json"
    config_path.write_text('{"sources": []}', encoding="utf-8")
    monkeypatch.chdir(script_path.parent)
    monkeypatch.setenv("ENERGY_FEED_CONFIG", str(config_path))
    monkeypatch.setattr(
        sys,
        "path",
        [entry for entry in sys.path if Path(entry or str(project_root)).resolve() != project_root],
    )

    runpy.run_path(str(script_path), run_name="__main__")

    assert str(project_root) in sys.path
    from scripts.ingest import upload_to_blob_storage

    assert callable(upload_to_blob_storage)